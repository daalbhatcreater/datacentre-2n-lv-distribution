"""Run the design and produce every deliverable.

    python build.py

results/load_schedule.xlsx   load schedule, breaker schedule, busbars, metering, BOM
results/ats_*.png            transfer sequence timelines
results/ats_*_events.csv     event logs of the same runs
drawings/sld.drawio, .png    single-line diagram
"""

import csv
import math
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from dcdesign import basis as B
from dcdesign import sld
from dcdesign.ats import max_step_kw, simulate
from dcdesign.sizing import current, design

ROOT = Path(__file__).parent
RESULTS = ROOT / "results"
DRAWINGS = ROOT / "drawings"
HEAD = PatternFill("solid", fgColor="1F3864")


def sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    for c, h in enumerate(header, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = (widths or {}).get(c, 14)
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=round(v, 2) if isinstance(v, float) else v)
    ws.freeze_panes = "A2"
    return ws


def load_schedule_rows(d):
    ups, dem = d["ups"], d["demand"]
    rows = []
    for case in ("normal", "failure"):
        u = ups[case]
        rows.append([case, "UPS-A input (IT load "
                     f"{u['load_kw']:.0f} kW at {100 * u['loading']:.0f} % UPS loading)",
                     "MSB-A", 1, u["input_kw"], B.UPS_INPUT_PF, u["input_kw"],
                     u["input_kva"], current(u["input_kva"])])
        share = 0.5 if case == "normal" else 1.0
        for m in d["mech"]:
            kw = m["kw"] * share
            kva = kw / m["pf"]
            rows.append([case, m["name"], m["board"], m["qty"] * share, m["kw_each"],
                         m["pf"], kw, kva, current(kva)])
        t = dem[case]
        rows.append([case, "TOTAL at MSB-A (vector sum)", "", "", "", round(t["pf"], 3),
                     t["kw"], t["kva"], t["current_a"]])
    return rows


def bom_rows(d):
    s, ups, bat, pdu = d["sources"], d["ups"], d["battery"], d["pdus"]
    rows = [
        ("Transformer", f"{s['transformer_kva']} kVA, 11/0.433 kV, Dyn11, uk 6.25 %, ONAN", 2),
        ("Diesel generator", f"{s['dg_kva']} kVA prime, 415 V, with AMF controller", 2),
        ("UPS module", f"{B.UPS_MODULE_KW:.0f} kW double conversion, unity pf", 2 * ups["modules"]),
        ("Battery block", "12 V VRLA high-rate, 10 min to 1.67 V/cell", 2 * bat["blocks"]),
        ("Battery string DC breaker", "DC MCCB, 2P, per string", 2 * bat["strings"]),
        ("Main LV switchboard", f"{d['busbars'][0]['rating_a']} A, "
         f"{d['busbars'][0]['icw_ka']} kA 1 s, Form 4b, IEC 61439-2", 2),
        ("UPS output switchboard", f"{d['busbars'][1]['rating_a']} A, Form 4b", 2),
        ("ATS", "630 A, 4-pole, open transition, mech + elec interlock", 4),
        ("PDU", f"{pdu['rating_kva']} kVA, 40-way, with branch circuit monitoring", 2 * pdu["per_side"]),
        ("Rack PDU", pdu["rack_pdu"], 2 * B.RACKS),
        ("Branch MCB", "16 A, 4-pole, C curve, 10 kA", 2 * B.RACKS),
    ]
    count = {}
    for b in d["breakers"]:
        key = (b.kind, f"{b.frame_a} A frame, 4-pole, Ir set {b.ir_a:.0f} A")
        count[key] = count.get(key, 0) + 2
    rows += [(k, spec, n) for (k, spec), n in sorted(count.items())]
    rows += [("Energy meter", m[1], 2 if "PDU" not in m[0] else 2 * pdu["per_side"])
             for m in d["metering"] if "Branch" not in m[1]]
    return rows


def write_workbook(d, path):
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "Load schedule", ["Case", "Load", "Board", "Qty", "kW each", "pf", "kW",
                                "kVA", "I (A)"],
          load_schedule_rows(d), {2: 48, 3: 12})
    sheet(wb, "Breakers (side A)",
          ["Tag", "Duty", "Type", "Frame (A)", "Ir (A)", "Design current (A)", "Ir loading"],
          [[b.tag, b.duty, b.kind, b.frame_a, b.ir_a, b.load_a, b.loading]
           for b in d["breakers"]], {2: 40})
    sheet(wb, "Busbars", ["Board", "Rating (A)", "Incomer (A)", "Sum of outgoing Ir (A)",
                          "RDF (IEC 61439-1)", "Diversified (A)", "Icw 1 s (kA)", "Form"],
          [list(b.values()) for b in d["busbars"]])
    s, ups, bat, pdu = d["sources"], d["ups"], d["battery"], d["pdus"]
    summary = [
        ("Transformer per side", f"{s['transformer_kva']} kVA",
         f"{100 * s['transformer_loading_normal']:.0f} % normal, "
         f"{100 * s['transformer_loading_failure']:.0f} % with B side lost"),
        ("Generator per side", f"{s['dg_kva']} kVA", f"{100 * s['dg_loading_failure']:.0f} % worst case"),
        ("UPS per side", f"{ups['modules']} x {B.UPS_MODULE_KW:.0f} kW",
         f"{100 * ups['normal']['loading']:.0f} % normal, {100 * ups['failure']['loading']:.0f} % failure"),
        ("Battery per side", f"{bat['strings']} strings x 40 blocks",
         f"{bat['w_per_cell_required']:.0f} W/cell needed, {bat['energy_kwh']:.0f} kWh at 1 MW"),
        ("PDUs per side", f"{pdu['per_side']} x {pdu['rating_kva']} kVA",
         f"{100 * pdu['loading_failure']:.0f} % with B side lost"),
        ("Design PUE (full IT load)", f"{d['pue']:.2f}", "transformer losses excluded"),
    ]
    sheet(wb, "Summary", ["Item", "Selected", "Note"], summary, {1: 28, 2: 24, 3: 48})
    sheet(wb, "Metering", ["Point", "Meter", "Purpose"], d["metering"], {1: 28, 2: 40, 3: 44})
    sheet(wb, "BOM (both sides)", ["Item", "Specification", "Qty"], bom_rows(d), {1: 26, 2: 60})
    wb.move_sheet("Summary", offset=-5)
    wb.save(path)


def plot_ats(res, path, title):
    t, S = res.t, res.series
    fig, ax = plt.subplots(4, 1, figsize=(10, 9), sharex=True)
    ax[0].plot(t, S["util_a"], label="Utility A")
    ax[0].plot(t, S["dg_a"], label="DG-A", ls="--")
    ax[0].plot(t, S["msb_a"], label="MSB-A bus", lw=2.2, alpha=0.6)
    ax[0].set_ylabel("Voltage (pu)")
    ax[0].set_ylim(-0.1, 1.2)
    ax[0].legend(loc="center right", fontsize=8)
    ax[1].plot(t, S["batt_kw"], color="tab:red", label="UPS-A battery power")
    ax[1].set_ylabel("Battery (kW)")
    ax1 = ax[1].twinx()
    ax1.plot(t, S["batt_pct"], color="tab:grey", ls=":", label="State of charge")
    ax1.set_ylabel("SOC (%)")
    ax1.set_ylim(0, 105)
    ax[1].legend(loc="upper right", fontsize=8)
    ax1.legend(loc="lower right", fontsize=8)
    ax[2].plot(t, S["msb_a_kw"], color="tab:green", label="MSB-A load")
    ax[2].plot(t, S["dg_kw"], color="tab:orange", ls="--", label="carried by DG-A")
    ax[2].set_ylabel("kW")
    ax[2].legend(loc="upper right", fontsize=8)
    ax[3].plot(t, S["it_a_kw"], label="IT load on A cords")
    ax[3].plot(t, S["it_b_kw"], label="IT load on B cords")
    ax[3].plot(t, S["it_a_kw"] + S["it_b_kw"], color="black", lw=1, label="total critical load")
    ax[3].set_ylabel("kW")
    ax[3].set_ylim(0, 1150)
    ax[3].legend(loc="center right", fontsize=8)
    ax[3].set_xlabel("Time (s)")
    for a in ax:
        a.grid(True, lw=0.3)
    fig.suptitle(title)
    fig.tight_layout()
    fig.savefig(path, dpi=130)
    plt.close(fig)


def write_events(res, path):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["t (s)", "event"])
        w.writerows(res.events)


def main():
    RESULTS.mkdir(exist_ok=True)
    DRAWINGS.mkdir(exist_ok=True)
    d = design()
    write_workbook(d, RESULTS / "load_schedule.xlsx")

    nodes, edges = sld.build(d)
    title = "Data centre LV distribution, 2N - single-line diagram"
    sld.render_png(nodes, edges, DRAWINGS / "sld.png", title)
    sld.write_drawio(nodes, edges, DRAWINGS / "sld.drawio", "SLD")

    scenarios = {
        "utility_fail": (dict(), "Utility A fails at 10 s, returns at 900 s; DG-A starts"),
        "dg_fail": (dict(dg_starts=False, utility_return_s=2400.0, t_end=3000.0),
                    "Utility A fails at 10 s and DG-A fails to start; utility returns at 2400 s"),
    }
    for key, (kw, title) in scenarios.items():
        res = simulate(**kw)
        plot_ats(res, RESULTS / f"ats_{key}.png", title)
        write_events(res, RESULTS / f"ats_{key}_events.csv")
        S = res.series
        print(f"\n{title}")
        for t, e in res.events:
            print(f"  {t:8.1f} s  {e}")
        print(f"  minimum battery SOC {S['batt_pct'].min():.1f} %, "
              f"largest DG load step {max_step_kw(S['dg_kw']):.0f} kW, "
              f"critical load never below {(S['it_a_kw'] + S['it_b_kw']).min():.0f} kW")

    s = d["sources"]
    print(f"\nTransformer {s['transformer_kva']} kVA "
          f"({100 * s['transformer_loading_failure']:.0f} % worst case), "
          f"DG {s['dg_kva']} kVA, UPS {d['ups']['modules']} x {B.UPS_MODULE_KW:.0f} kW, "
          f"{d['battery']['strings']} battery strings, PUE {d['pue']:.2f}")


if __name__ == "__main__":
    main()
